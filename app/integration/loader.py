"""Explicit column maps, bounded input and atomic CSV/XLSX import."""
import csv
import json
from pathlib import Path
from pydantic import ValidationError
from sqlalchemy.exc import IntegrityError
from app.services.dataset_loader import DatasetLoader, LoadReport
from app.core.exceptions import DataValidationError
from app.integration.registry import contracts

class TabularLoader(DatasetLoader):
    def __init__(self, repository, agent, *, fixture=False, store_names=None):
        self.repository, self.agent = repository, agent
        self.fixture, self.store_names = fixture, store_names

    def load(self, path:Path, column_mapping:dict[str,str])->LoadReport:
        self.validate_source(path)
        if path.stat().st_size>20_000_000: raise DataValidationError('Dataset exceeds 20 MB')
        if path.suffix.lower()=='.csv':
            with path.open(encoding='utf-8-sig',newline='') as stream:
                rows=list(csv.reader(stream))
        else:
            from openpyxl import load_workbook
            workbook=load_workbook(path,read_only=True,data_only=False)
            try:
                sheet=workbook.active
                if sheet.max_row>10001 or sheet.max_column>100: raise DataValidationError('Workbook too large')
                rows=[list(row) for row in sheet.iter_rows(values_only=True)]
            finally: workbook.close()
        if not rows or not rows[0]: raise DataValidationError('Empty dataset')
        header=rows[0]
        if len(header)!=len(set(header)) or any(not isinstance(v,str) or not v.strip() for v in header):
            raise DataValidationError('Invalid or duplicate column headers')
        if len(rows)>10001: raise DataValidationError('Import limited to 10000 rows')
        model,_,_=contracts(self.agent)
        if not column_mapping or not set(column_mapping)<=set(header): raise DataValidationError('Explicit source-to-field mapping required')
        fields=list(column_mapping.values())
        if len(fields)!=len(set(fields)) or not set(fields)<=set(model.model_fields): raise DataValidationError('Invalid target mapping')
        parsed=[]
        try:
            for number,values in enumerate(rows[1:],2):
                if len(values)!=len(header): raise ValueError('Ragged row')
                source=dict(zip(header,values)); record={}
                for col,field in column_mapping.items():
                    value=source[col]
                    if isinstance(value,str) and value.startswith('='):raise ValueError('Formula input is unsupported')
                    if value in ('',None): continue
                    if isinstance(value,str) and value.lstrip().startswith(('[','{')):value=json.loads(value)
                    record[field]=value
                # Provenance and fixture label are importer-controlled, never taken from the sheet.
                record['source']=path.name
                record['fixture']=self.fixture
                parsed.append(model.model_validate(record))
            if not parsed: raise ValueError('No data rows')
            loaded=self.repository.add_records(self.agent,parsed,store_names=self.store_names)
        except (ValueError,ValidationError,IntegrityError) as exc:
            raise DataValidationError('Dataset rejected; verify required fields, dates, values and unique identities') from None
        return LoadReport(source_name=path.name,rows_seen=len(parsed),rows_loaded=loaded,column_mapping=column_mapping)
