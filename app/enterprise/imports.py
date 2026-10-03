"""Trusted local workbook import; atomic, bounded and idempotent by file hash."""
import argparse
from collections import Counter
from datetime import date, datetime, time
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path
import re
from openpyxl import load_workbook
from sqlalchemy import insert, select
from app.models.database import Base
from app.integration.repository import StoreRow
from app.services.database_service import Database
from app.conversation.taxonomy import classify
from app.enterprise.data import ImportBatch, StoreMetadata, Product, Sale, CompetitorSale

SALES_COLUMNS = 'txn_id order_id sale_date sale_time day_of_week time_slot store_id store_location city zone sales_channel sku_id style_code style_name category gender color size quantity unit_price_inr discount_pct net_amount_inr payment_mode customer_type fulfillment_type delivery_partner delivery_status delivery_status_detail promised_delivery_date actual_delivery_date delay_days_vs_promise benchmark_competitor competitor_price_inr competitor_promo_active price_index_vs_competitor price_position'.split()
COMPETITOR_COLUMNS = 'week_start store_id city category competitor competitor_promo_active competitor_avg_discount_pct competitor_avg_price_inr competitor_units_sold_est competitor_revenue_est_inr our_units_sold our_revenue_inr our_share_of_tracked_market_pct'.split()
STORE_COLUMNS = 'store_id city store_location zone store_format delivery_tier store_area_sqft demand_index primary_competitor competitors_tracked'.split()


def sheets(path, required):
    path=Path(path)
    if not path.is_file() or path.suffix.lower()!='.xlsx':raise ValueError('Supply an existing XLSX workbook.')
    if path.stat().st_size>100_000_000:raise ValueError('Workbook exceeds 100 MB.')
    book=load_workbook(path,read_only=True,data_only=False)
    try:
        if not set(required)<=set(book.sheetnames) or 'Read Me' not in book.sheetnames:raise ValueError('Required workbook sheets are missing.')
        note=' '.join(str(v) for row in book['Read Me'].values for v in row if v is not None)
        if 'synthetic' not in note.lower():raise ValueError('This POC importer requires a source explicitly labelled synthetic.')
        result={}
        for name,columns in required.items():
            iterator=iter(book[name].values);header=next(iterator)
            if len(header)!=len(set(header)) or not set(columns)<=set(header):raise ValueError('Missing or duplicate columns in '+name)
            rows=[]
            for i,values in enumerate(iterator,2):
                if i>200001:raise ValueError('Sheet exceeds 200,000 rows.')
                if not any(v is not None for v in values):continue
                if any(isinstance(v,str) and v.startswith('=') for v in values):raise ValueError('Formula cells are unsupported in '+name)
                row={k:(v.isoformat() if isinstance(v,(date,datetime,time)) else v) for k,v in zip(header,values)}
                rows.append(row)
            if not rows:raise ValueError('Empty sheet: '+name)
            result[name]=rows
        return result
    finally:book.close()


def number(value, *, minimum=None, maximum=None):
    try:n=Decimal(str(value))
    except Exception:raise ValueError('Invalid numeric value') from None
    if not n.is_finite() or minimum is not None and n<minimum or maximum is not None and n>maximum:raise ValueError('Numeric value outside permitted range')
    return n


def integer(value, minimum=0):
    n=number(value,minimum=minimum)
    if n!=int(n):raise ValueError('Expected whole number')
    return int(n)


def paise(value):return int((number(value,minimum=0)*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))


def identifier(value):
    if not isinstance(value,str) or not re.fullmatch(r'[\w.-]{1,128}',value):raise ValueError('Invalid record identifier')
    return value


def day(value):
    try:return date.fromisoformat(str(value)[:10]).isoformat()
    except ValueError:raise ValueError('Invalid source date') from None


def report_for(path,domain,data,expected):
    return {'source':Path(path).name,'sha256':hashlib.sha256(Path(path).read_bytes()).hexdigest(),'domain':domain,'synthetic':True,
        'counts':{k:len(v) for k,v in data.items()},'missing':{k:{c:sum(r.get(c) is None for r in rows) for c in rows[0]} for k,rows in data.items()},
        'warnings':[f'{k}: expected approximately {n} rows; observed {len(data[k])}.' for k,n in expected.items() if len(data[k])!=n]}


def save_stores(session, records):
    for r in records:
        sid=identifier(r['store_id']);name=r['city']
        if not isinstance(name,str) or not name.strip():raise ValueError('Store city is required')
        existing=session.get(StoreRow,sid)
        if existing and existing.name!=name:raise ValueError('Store identity conflicts with loaded data')
        if not existing:session.add(StoreRow(id=sid,name=name));session.flush()
        existing_meta=session.get(StoreMetadata,sid)
        if not existing_meta:session.add(StoreMetadata(store_id=sid,payload=r))


def import_sales(database,path):
    Base.metadata.create_all(database.engine)
    digest='sales:'+hashlib.sha256(Path(path).read_bytes()).hexdigest()
    with database.session() as session:
        prior=session.get(ImportBatch,digest)
        if prior:return {**prior.report,'already_imported':True}
    data=sheets(path,{'Sales_Transactions':SALES_COLUMNS,'Competitor_Sales':COMPETITOR_COLUMNS,'Store_Master':STORE_COLUMNS})
    report=report_for(path,'sales',data,{'Sales_Transactions':75000,'Competitor_Sales':24804,'Store_Master':20})
    stores=data['Store_Master'];ids={identifier(r['store_id']) for r in stores}
    if len(ids)!=len(stores):raise ValueError('Duplicate stores')
    products={};sales=[];competitors=[];seen=set()
    for r in data['Sales_Transactions']:
        tid=identifier(r['txn_id']);sku=identifier(r['sku_id']);identifier(r['order_id'])
        if tid in seen or r['store_id'] not in ids:raise ValueError('Duplicate transaction or unknown store')
        seen.add(tid);when=day(r['sale_date']);quantity=integer(r['quantity'],1)
        price=number(r['unit_price_inr'],minimum=0);discount=number(r['discount_pct'],minimum=0,maximum=100)
        net=paise(r['net_amount_inr'])
        # Source net amounts are rounded to whole rupees; preserve them, within INR 0.50 of the formula.
        if abs(net-paise(quantity*price*(1-discount/100)))>50:raise ValueError('Net revenue does not reconcile with quantity, price and discount')
        for field in ['promised_delivery_date','actual_delivery_date']:
            if r[field] is not None and day(r[field])<when:raise ValueError('Delivery date precedes sale')
        product={k:r[k] for k in ['sku_id','style_code','style_name','category','gender','color','size']}
        product.update(classify(r['category']))
        if sku in products and products[sku]!=product:raise ValueError('Conflicting product attributes')
        products[sku]=product
        sales.append(dict(txn_id=tid,order_id=r['order_id'],store_id=r['store_id'],sku_id=sku,day=when,category=product['normalized_category'] or r['category'],channel=r['sales_channel'],quantity=quantity,net_paise=net,payload=r))
    seen=set()
    for r in data['Competitor_Sales']:
        when=day(r['week_start']);identity=(when,r['store_id'],r['category'],r['competitor'])
        if identity in seen or r['store_id'] not in ids:raise ValueError('Duplicate competitor record or unknown store')
        seen.add(identity)
        number(r['our_share_of_tracked_market_pct'],minimum=0,maximum=100)
        for field in ['competitor_avg_discount_pct'] :number(r[field],minimum=0,maximum=100)
        for field in ['competitor_avg_price_inr','our_revenue_inr','our_units_sold']:number(r[field],minimum=0)
        competitors.append(dict(week=when,store_id=r['store_id'],category=classify(r['category'])['normalized_category'] or r['category'],competitor=r['competitor'],units=integer(r['competitor_units_sold_est']),revenue_paise=paise(r['competitor_revenue_est_inr']),payload=r))
    report.update(orders=len({r['order_id'] for r in sales}),skus=len(products),date_range=[min(r['day'] for r in sales),max(r['day'] for r in sales)],net_revenue_inr=str(Decimal(sum(r['net_paise'] for r in sales))/100))
    with database.session() as session:
        save_stores(session,stores)
        for sku,payload in products.items():
            existing=session.get(Product,sku)
            if existing and existing.payload!=payload:raise ValueError('Product conflicts with imported version')
            if not existing:session.add(Product(sku_id=sku,payload=payload))
        session.flush()
        for model,rows in [(Sale,sales),(CompetitorSale,competitors)]:
            for start in range(0,len(rows),1000):session.execute(insert(model),rows[start:start+1000])
        session.add(ImportBatch(id=digest,domain='sales',report=report))
    return report


def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--file',type=Path,required=True);parser.add_argument('--database',required=True)
    args=parser.parse_args();db=Database(args.database)
    try:print(json.dumps(import_sales(db,args.file),indent=2))
    finally:db.close()


if __name__=='__main__':main()
